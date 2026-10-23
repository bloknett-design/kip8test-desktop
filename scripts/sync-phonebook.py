#!/usr/bin/env python3
"""
Синхронизация телефонного справочника с Google Sheets.

Источник (таблица пользователя, Task 542; ссылка обновлена Task 545):
  https://docs.google.com/spreadsheets/d/1rFbDn-jNRSoELTnsiaKjNvatSqIhVvbS/edit
  лист: "Лист1" (плоская таблица из 3 колонок: Должность | Ф.И.О. | № тел.)

Скрипт работает по тому же принципу, что и scripts/sync-cables.py
(раздел «Кабельный журнал»):

  1. Скачивает XLSX-экспорт напрямую из Google Sheets через
     export?format=xlsx. Google отдаёт файл без OAuth, если таблица
     доступна «у кого есть ссылка».
  2. Парсит лист "Лист1" — строка 1 это заголовки колонок
     (Column1/Column2/Column3), данные со 2-й строки.
  3. Сохраняет результат в data/phonebook.json.

ПРАВИЛА РАЗБОРА (плоская таблица БЕЗ форматирования — аквамари-
новые заливки исходного XLS при импорте в Google Sheets потерялись,
поэтому структура определяется только по наполнению строк):

  Строка «A-only» (заполнена только колонка A):
    - начинается с «ООО» → ЗАГОЛОВОК ОРГАНИЗАЦИИ (новая секция);
    - иначе            → ЗАГОЛОВОК ПОДРАЗДЕЛЕНИЯ (группа) внутри
                         текущей организации.
  Любая строка, где заполнена колонка B и/или C → КОНТАКТ текущей
  организации/подразделения: position=A, name=B, phone=C (пустые —
  пустой строкой). Пустая A — строка-продолжение (position="").
  Строки ДО первого заголовка организации относятся к организации
  ООО ПО «ТОКЕМ» (DEFAULT_HEADER — как parsing_rules.default_header
  исходного XLS; постоянство ключа сохраняет ID избранного).

  Подгруппы третьего уровня (subgroup) плоской таблицей не выра-
  жаются — всегда null. Схема JSON сохранена для совместимости.

Переменные окружения:
  PHONEBOOK_SPREADSHEET_ID — ID Google Sheets
      (по умолчанию 1rFbDn-jNRSoELTnsiaKjNvatSqIhVvbS — Task 545)
  PHONEBOOK_SHEET_NAME — имя листа (по умолчанию "Лист1")

Секреты НЕ требуются — таблица доступна «у кого есть ссылка».

Если нет интернета или API недоступен — существующий
data/phonebook.json остаётся без изменений (PWA продолжает работать
с последними закоммиченными данными).
"""

import json
import sys
from pathlib import Path

import requests

try:
    import openpyxl
except ImportError:
    print('[phonebook] openpyxl не установлен', flush=True)
    sys.exit(1)


# ============================================================
# Настройки Google Sheets
# ============================================================
DEFAULT_SPREADSHEET_ID = '1rFbDn-jNRSoELTnsiaKjNvatSqIhVvbS'
DEFAULT_SHEET_NAME = 'Лист1'
SHEET_URL = ('https://docs.google.com/spreadsheets/d/'
             + DEFAULT_SPREADSHEET_ID + '/edit')

# Организация по умолчанию: строки до первого явного заголовка «ООО …».
# Совпадает с parsing_rules.default_header исходного XLS — ключ секции
# (и ID избранного header|group|…) остаются стабильными.
DEFAULT_HEADER = 'ООО ПО «ТОКЕМ»'
TITLE = 'Телефонный справочник ООО ПО «ТОКЕМ»'

PROJECT_ROOT = Path(__file__).resolve().parent.parent
JSON_OUT = PROJECT_ROOT / 'data' / 'phonebook.json'

PARSING_RULES = {
    'source': 'Google Sheets (плоская таблица 3 колонок: Должность / '
              'Ф.И.О. / № тел.) — Task 542',
    'org_header': 'только колонка A + значение начинается с «ООО»',
    'group_header': 'только колонка A (не «ООО…») — подразделение '
                    'текущей организации',
    'contact_row': 'заполнена колонка B и/или C: position=A, name=B, '
                   'phone=C; пустая A — строка-продолжение (position="")',
    'default_header': DEFAULT_HEADER,
    'subgroup': 'плоской таблицей не выражается — всегда null',
}


def log(msg):
    print(f'[phonebook] {msg}', flush=True)


def download_sheet():
    """Скачивает XLSX-экспорт книги Google Sheets (без OAuth)."""
    ssid = __import__('os').environ.get(
        'PHONEBOOK_SPREADSHEET_ID', DEFAULT_SPREADSHEET_ID)
    url = (f'https://docs.google.com/spreadsheets/d/{ssid}'
           '/export?format=xlsx')
    log(f'Скачиваю: {url}')
    resp = requests.get(url, timeout=60)
    resp.raise_for_status()
    out = Path('/tmp/phonebook_download.xlsx')
    out.write_bytes(resp.content)
    log(f'Скачано {len(resp.content)} байт → {out}')
    return out


def parse_sheet(xlsx_path):
    """Разбирает плоскую таблицу в структуру phonebook.json."""
    sheet_name = __import__('os').environ.get(
        'PHONEBOOK_SHEET_NAME', DEFAULT_SHEET_NAME)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        log(f'ОШИБКА: лист «{sheet_name}» не найден; доступно: '
            f'{wb.sheetnames}')
        return None
    ws = wb[sheet_name]

    sections = {}
    order = []            # порядок организаций (ключи секций)
    current_org = DEFAULT_HEADER
    current_group = None

    def ensure_org(header):
        nonlocal current_org
        current_org = header
        if header not in sections:
            sections[header] = []
            order.append(header)

    ensure_org(DEFAULT_HEADER)

    n_rows = 0
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=3):
        a = (str(r[0].value).strip() if r[0].value is not None else '')
        b = (str(r[1].value).strip() if r[1].value is not None else '')
        c = (str(r[2].value).strip() if r[2].value is not None else '')
        if not a and not b and not c:
            continue          # пустая строка — пропускаем
        n_rows += 1
        if a and not b and not c:
            # Строка-заголовок
            if a.startswith('ООО'):
                ensure_org(a)
                current_group = None
                log(f'Организация: {a}')
            else:
                current_group = a
                log(f'  Подразделение: {a}')
            continue
        # Строка контакта (B и/или C заполнены)
        sections[current_org].append({
            'group': current_group,
            'subgroup': None,
            'position': a,
            'name': b,
            'phone': c,
        })

    # Убрать организации без единой строки (заголовок в конце и т.п.)
    for h in list(order):
        if not sections[h]:
            del sections[h]
            order.remove(h)

    total_records = sum(len(v) for v in sections.values())
    log(f'Строк данных: {n_rows}; организаций: {len(order)}; '
        f'записей: {total_records}')
    return {
        'title': TITLE,
        'source': SHEET_URL,
        'total_records': total_records,
        'total_companies': len(order),
        'parsing_rules': PARSING_RULES,
        'sections': {h: sections[h] for h in order},
    }


def main():
    try:
        xlsx = download_sheet()
    except Exception as e:
        log(f'Скачать не удалось ({e}) — оставляю существующий '
            f'{JSON_OUT.name}')
        return 0
    data = parse_sheet(xlsx)
    if data is None:
        log('Разбор не удался — оставляю существующий JSON')
        return 1
    if not data['sections']:
        log('Пустой результат — оставляю существующий JSON')
        return 1

    JSON_OUT.parent.mkdir(parents=True, exist_ok=True)
    old = ''
    if JSON_OUT.exists():
        old = JSON_OUT.read_text(encoding='utf-8')
    new = json.dumps(data, ensure_ascii=False, indent=2) + '\n'
    if new == old:
        log('Изменений нет')
        return 0
    JSON_OUT.write_text(new, encoding='utf-8')
    log(f'Обновлён {JSON_OUT} '
        f'({data["total_records"]} записей, '
        f'{data["total_companies"]} организаций)')
    return 0


if __name__ == '__main__':
    sys.exit(main())
